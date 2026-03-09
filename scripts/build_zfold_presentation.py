from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / ".vendor"))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = ROOT / "docs" / "zfold-presentation.ja.pptx"


def rgb(hex_value: str) -> RGBColor:
    hex_value = hex_value.replace("#", "")
    return RGBColor.from_string(hex_value)


BG = rgb("F6F1E8")
PAPER = rgb("FFFCF7")
INK = rgb("1F2937")
MUTED = rgb("5B6472")
TEAL = rgb("0F766E")
TEAL_SOFT = rgb("D8F0EA")
ORANGE = rgb("D97706")
ORANGE_SOFT = rgb("FCE7C2")
NAVY = rgb("19486A")
NAVY_SOFT = rgb("DDEAF5")
RED = rgb("B45309")
RED_SOFT = rgb("FDE7D9")
GREEN = rgb("3F7D58")
GREEN_SOFT = rgb("DCEFD9")
WHITE = rgb("FFFFFF")
LINE = rgb("D8D0C2")


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def blank_slide():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    return slide


def add_rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape_type = (
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        if radius
        else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    )
    shape = slide.shapes.add_shape(shape_type, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill if line is None else line
    shape.line.width = Pt(1)
    return shape


def add_circle(slide, x, y, d, fill, line=None):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x, y, d, d)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill if line is None else line
    shape.line.width = Pt(1)
    return shape


def add_text(
    slide,
    x,
    y,
    w,
    h,
    text,
    size=18,
    color=INK,
    bold=False,
    font="Yu Gothic UI",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(4)
    tf.margin_right = Pt(4)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(
    slide,
    x,
    y,
    w,
    h,
    bullets,
    size=18,
    color=INK,
    font="Yu Gothic UI",
):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(4)
    tf.margin_right = Pt(4)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    first = True
    for bullet in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = f"• {bullet}"
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(7)
        for run in p.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.color.rgb = color
    return box


def add_title(slide, overline, title, subtitle=None):
    add_rect(slide, Inches(0.55), Inches(0.42), Inches(1.0), Inches(0.08), TEAL)
    add_text(slide, Inches(0.65), Inches(0.55), Inches(4.8), Inches(0.35), overline, 12, TEAL, True)
    add_text(slide, Inches(0.65), Inches(0.9), Inches(11.5), Inches(0.7), title, 28, INK, True)
    if subtitle:
        add_text(slide, Inches(0.68), Inches(1.55), Inches(11.2), Inches(0.55), subtitle, 14, MUTED)


def add_footer(slide, page):
    add_text(
        slide,
        Inches(12.55),
        Inches(7.02),
        Inches(0.45),
        Inches(0.22),
        str(page),
        10,
        MUTED,
        False,
        align=PP_ALIGN.RIGHT,
    )


def add_card(slide, x, y, w, h, title, body_lines, accent, soft_fill):
    add_rect(slide, x, y, w, h, PAPER, LINE, radius=True)
    add_rect(slide, x, y, Inches(0.10), h, accent, accent)
    add_text(slide, x + Inches(0.22), y + Inches(0.15), w - Inches(0.4), Inches(0.35), title, 18, INK, True)
    add_rect(slide, x + Inches(0.22), y + Inches(0.56), w - Inches(0.38), Inches(0.02), soft_fill, soft_fill)
    add_bullets(slide, x + Inches(0.22), y + Inches(0.72), w - Inches(0.34), h - Inches(0.85), body_lines, 15, MUTED)


def add_pill(slide, x, y, w, h, text, fill, color):
    shape = add_rect(slide, x, y, w, h, fill, fill, radius=True)
    add_text(slide, x, y + Pt(1), w, h, text, 12, color, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    return shape


def add_arrow(slide, x1, y1, x2, y2, color):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = color
    line.line.width = Pt(2)
    line.line.end_arrowhead = True
    return line


def add_file_stack(slide, x, y, labels, fill):
    for idx, label in enumerate(labels):
        off = Inches(0.16) * idx
        add_rect(slide, x + off, y + off, Inches(1.3), Inches(0.58), fill, WHITE, radius=True)
        add_text(
            slide,
            x + off + Inches(0.08),
            y + off + Inches(0.10),
            Inches(1.1),
            Inches(0.24),
            label,
            12,
            INK,
            True,
        )


def add_zip_vs_zfold(slide):
    left_x = Inches(0.75)
    right_x = Inches(6.9)
    top_y = Inches(1.75)

    add_card(
        slide,
        Inches(0.55),
        Inches(1.45),
        Inches(5.8),
        Inches(4.9),
        "ZIP: ファイルごとに個別圧縮",
        [
            "取り出しやすい",
            "汎用性が高い",
            "ファイルをまたぐ共通パターンは活かしにくい",
        ],
        ORANGE,
        ORANGE_SOFT,
    )
    add_card(
        slide,
        Inches(6.95),
        Inches(1.45),
        Inches(5.8),
        Inches(4.9),
        "zfold: 複数ファイルをチャンク単位で圧縮",
        [
            "小さい類似ファイル群で圧縮率が上がりやすい",
            "同じテンプレートやキー構造を拾いやすい",
            "一部展開時はチャンク全体を読む",
        ],
        TEAL,
        TEAL_SOFT,
    )

    add_file_stack(slide, left_x, top_y, ["a.html", "b.html", "c.json"], PAPER)
    add_arrow(slide, Inches(2.55), Inches(2.55), Inches(3.2), Inches(2.55), ORANGE)
    add_rect(slide, Inches(3.25), Inches(2.15), Inches(1.2), Inches(0.82), ORANGE_SOFT, ORANGE, radius=True)
    add_text(slide, Inches(3.33), Inches(2.38), Inches(1.05), Inches(0.2), "別々に圧縮", 12, ORANGE, True, align=PP_ALIGN.CENTER)
    add_arrow(slide, Inches(4.5), Inches(2.55), Inches(5.25), Inches(2.55), ORANGE)
    add_rect(slide, Inches(5.3), Inches(2.15), Inches(0.78), Inches(0.82), ORANGE, ORANGE, radius=True)
    add_text(slide, Inches(5.42), Inches(2.37), Inches(0.5), Inches(0.2), "ZIP", 14, WHITE, True, align=PP_ALIGN.CENTER)

    add_file_stack(slide, right_x, top_y, ["a.html", "b.html", "c.json"], PAPER)
    add_arrow(slide, Inches(8.7), Inches(2.55), Inches(9.35), Inches(2.55), TEAL)
    add_rect(slide, Inches(9.4), Inches(2.05), Inches(1.35), Inches(1.02), TEAL_SOFT, TEAL, radius=True)
    add_text(slide, Inches(9.53), Inches(2.18), Inches(1.05), Inches(0.2), "chunk", 13, TEAL, True, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(9.48), Inches(2.45), Inches(1.14), Inches(0.28), "まとめて圧縮", 12, MUTED, False, align=PP_ALIGN.CENTER)
    add_arrow(slide, Inches(10.8), Inches(2.55), Inches(11.45), Inches(2.55), TEAL)
    add_rect(slide, Inches(11.5), Inches(2.15), Inches(0.78), Inches(0.82), TEAL, TEAL, radius=True)
    add_text(slide, Inches(11.58), Inches(2.37), Inches(0.62), Inches(0.2), ".zpk", 14, WHITE, True, align=PP_ALIGN.CENTER)

    add_pill(slide, Inches(4.05), Inches(5.7), Inches(5.2), Inches(0.38), "差が出るのは「似た小ファイルが大量にある」ケース", NAVY_SOFT, NAVY)


def add_chunk_scale(slide):
    add_rect(slide, Inches(0.9), Inches(2.1), Inches(11.4), Inches(0.16), LINE)
    for x in [1.4, 4.25, 7.1, 9.95]:
        add_rect(slide, Inches(x), Inches(1.98), Inches(0.03), Inches(0.42), NAVY)
    add_circle(slide, Inches(5.92), Inches(1.79), Inches(0.46), ORANGE)
    add_text(slide, Inches(5.75), Inches(2.32), Inches(0.8), Inches(0.22), "既定値", 12, ORANGE, True, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(0.88), Inches(2.46), Inches(2.4), Inches(0.32), "小さい chunk-size", 15, TEAL, True)
    add_text(slide, Inches(9.6), Inches(2.46), Inches(2.2), Inches(0.32), "大きい chunk-size", 15, ORANGE, True, align=PP_ALIGN.RIGHT)
    add_card(
        slide,
        Inches(0.8),
        Inches(3.05),
        Inches(3.8),
        Inches(2.35),
        "小さめにすると",
        [
            "メモリを抑えやすい",
            "一部展開の無駄が減りやすい",
            "圧縮率は少し下がることがある",
        ],
        TEAL,
        TEAL_SOFT,
    )
    add_card(
        slide,
        Inches(4.78),
        Inches(3.05),
        Inches(3.8),
        Inches(2.35),
        "大きめにすると",
        [
            "圧縮率が少し良くなることがある",
            "圧縮時メモリが増える",
            "部分展開では読む量が増えやすい",
        ],
        ORANGE,
        ORANGE_SOFT,
    )
    add_card(
        slide,
        Inches(8.76),
        Inches(3.05),
        Inches(3.75),
        Inches(2.35),
        "低メモリ環境の目安",
        [
            "まず既定値を試す",
            "厳しければ --threads 1",
            "さらに --chunk-size を小さくする",
        ],
        NAVY,
        NAVY_SOFT,
    )
    add_rect(slide, Inches(3.15), Inches(6.0), Inches(7.1), Inches(0.6), PAPER, LINE, radius=True)
    add_text(
        slide,
        Inches(3.35),
        Inches(6.18),
        Inches(6.7),
        Inches(0.25),
        "例: zfold pack ./dataset -o dataset.zpk --threads 1 --chunk-size 8388608",
        15,
        INK,
        True,
    )


def add_dictionary_flow(slide):
    step_y = Inches(2.0)
    widths = [Inches(2.3), Inches(2.0), Inches(2.2), Inches(2.4)]
    xs = [Inches(0.75), Inches(3.45), Inches(5.9), Inches(8.6)]
    titles = ["学習用サンプル", "train-dict", "辞書ファイル app.dict", "pack --dict で圧縮"]
    fills = [PAPER, ORANGE_SOFT, NAVY_SOFT, TEAL_SOFT]
    lines = [LINE, ORANGE, NAVY, TEAL]
    bodies = [
        "似た HTML / JSON / 設定ファイル",
        "共通パターンを抽出",
        "圧縮のヒント集",
        "辞書もアーカイブに埋め込まれる",
    ]
    for x, w, title, fill, line, body in zip(xs, widths, titles, fills, lines, bodies):
        add_rect(slide, x, step_y, w, Inches(1.35), fill, line, radius=True)
        add_text(slide, x + Inches(0.12), step_y + Inches(0.17), w - Inches(0.22), Inches(0.28), title, 16, INK, True)
        add_text(slide, x + Inches(0.12), step_y + Inches(0.6), w - Inches(0.22), Inches(0.42), body, 13, MUTED)
    add_arrow(slide, Inches(3.1), Inches(2.67), Inches(3.42), Inches(2.67), ORANGE)
    add_arrow(slide, Inches(5.52), Inches(2.67), Inches(5.87), Inches(2.67), NAVY)
    add_arrow(slide, Inches(8.22), Inches(2.67), Inches(8.56), Inches(2.67), TEAL)

    add_card(
        slide,
        Inches(0.78),
        Inches(4.15),
        Inches(3.9),
        Inches(1.8),
        "辞書のメリット",
        [
            "小さい類似ファイルで効きやすい",
            "同じ種類のデータを何度も固める運用に向く",
        ],
        TEAL,
        TEAL_SOFT,
    )
    add_card(
        slide,
        Inches(4.87),
        Inches(4.15),
        Inches(3.9),
        Inches(1.8),
        "効きやすい例",
        [
            "HTML テンプレート",
            "共通キー構造の JSON",
            "INI / TOML / YAML / ログ",
        ],
        NAVY,
        NAVY_SOFT,
    )
    add_card(
        slide,
        Inches(8.96),
        Inches(4.15),
        Inches(3.6),
        Inches(1.8),
        "効きにくい例",
        [
            "JPEG / PNG / MP4",
            "ZIP など圧縮済みデータ",
            "共通性の少ないバイナリ",
        ],
        ORANGE,
        ORANGE_SOFT,
    )


def add_commands_flow(slide):
    cards = [
        (Inches(0.8), NAVY_SOFT, NAVY, "1. 任意で辞書を作る", "zfold train-dict ./snap1 ./snap2 -o app.dict --mode text"),
        (Inches(3.95), TEAL_SOFT, TEAL, "2. pack で固める", "zfold pack ./project -o project.zpk --dict app.dict"),
        (Inches(7.1), PAPER, ORANGE, "3. verify / list で確認", "zfold verify project.zpk\nzfold list project.zpk"),
        (Inches(10.25), ORANGE_SOFT, ORANGE, "4. extract で取り出す", "zfold extract project.zpk -d ./out --prefix \"docs/\""),
    ]
    for x, fill, line, title, body in cards:
        add_rect(slide, x, Inches(2.05), Inches(2.3), Inches(2.4), fill, line, radius=True)
        add_text(slide, x + Inches(0.12), Inches(2.22), Inches(2.06), Inches(0.45), title, 16, INK, True)
        add_text(slide, x + Inches(0.12), Inches(2.92), Inches(2.04), Inches(1.1), body, 12, MUTED, False, font="Consolas")
    add_arrow(slide, Inches(3.18), Inches(3.22), Inches(3.9), Inches(3.22), NAVY)
    add_arrow(slide, Inches(6.33), Inches(3.22), Inches(7.05), Inches(3.22), TEAL)
    add_arrow(slide, Inches(9.48), Inches(3.22), Inches(10.2), Inches(3.22), ORANGE)

    add_rect(slide, Inches(0.85), Inches(5.1), Inches(5.95), Inches(1.25), PAPER, LINE, radius=True)
    add_text(slide, Inches(1.05), Inches(5.28), Inches(5.55), Inches(0.24), "暗号化したいとき", 17, INK, True)
    add_bullets(
        slide,
        Inches(1.02),
        Inches(5.58),
        Inches(5.55),
        Inches(0.56),
        ["--password-prompt か --password-file が扱いやすい", "チャンク、辞書、index が暗号化される"],
        14,
        MUTED,
    )
    add_rect(slide, Inches(6.95), Inches(5.1), Inches(5.55), Inches(1.25), PAPER, LINE, radius=True)
    add_text(slide, Inches(7.15), Inches(5.28), Inches(5.15), Inches(0.24), "運用で押さえる点", 17, INK, True)
    add_bullets(
        slide,
        Inches(7.12),
        Inches(5.58),
        Inches(5.15),
        Inches(0.56),
        ["バックアップ用途では pack 後に verify を回す", "辞書はアーカイブに埋め込まれるので別配布不要"],
        14,
        MUTED,
    )


def add_intro_slide():
    slide = blank_slide()
    add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(7.5), BG, BG)
    add_rect(slide, Inches(0), Inches(0), Inches(4.6), Inches(7.5), NAVY, NAVY)
    add_rect(slide, Inches(4.2), Inches(0.75), Inches(8.7), Inches(6.0), PAPER, PAPER, radius=True)
    add_circle(slide, Inches(9.6), Inches(0.6), Inches(2.0), ORANGE_SOFT, ORANGE_SOFT)
    add_circle(slide, Inches(11.1), Inches(1.5), Inches(1.05), TEAL_SOFT, TEAL_SOFT)
    add_text(slide, Inches(0.62), Inches(1.0), Inches(3.3), Inches(0.32), "GUIDE DECK", 12, TEAL_SOFT, True, font="Bahnschrift")
    add_text(slide, Inches(0.62), Inches(1.45), Inches(3.4), Inches(1.55), "zfold\n入門", 28, WHITE, True)
    add_text(slide, Inches(0.64), Inches(3.1), Inches(3.25), Inches(1.2), "ZIP と何が違うのか\n辞書はいつ効くのか\nどんな場面で使い分けるのか", 17, TEAL_SOFT)
    add_pill(slide, Inches(4.75), Inches(1.15), Inches(1.8), Inches(0.36), ".zpk archive", NAVY_SOFT, NAVY)
    add_text(slide, Inches(4.8), Inches(1.75), Inches(7.3), Inches(0.58), "似た小ファイル群を、\nきれいに小さくまとめるためのアーカイバ", 25, INK, True)
    add_text(slide, Inches(4.82), Inches(3.0), Inches(7.2), Inches(0.82), "solid archive 的な圧縮と zstd、さらに辞書圧縮を組み合わせて、\nWeb 配布物、設定ファイル群、スナップショットの保存に向く設計です。", 16, MUTED)
    add_card(
        slide,
        Inches(4.75),
        Inches(4.55),
        Inches(2.35),
        Inches(1.4),
        "強い場面",
        ["類似した小ファイルが多い", "圧縮率を重視したい"],
        TEAL,
        TEAL_SOFT,
    )
    add_card(
        slide,
        Inches(7.38),
        Inches(4.55),
        Inches(2.35),
        Inches(1.4),
        "注意点",
        ["一部展開はチャンク単位", "設定次第でメモリを使う"],
        ORANGE,
        ORANGE_SOFT,
    )
    add_card(
        slide,
        Inches(10.0),
        Inches(4.55),
        Inches(2.35),
        Inches(1.4),
        "辞書対応",
        ["小さい類似データに効く", "辞書もアーカイブへ同梱"],
        NAVY,
        NAVY_SOFT,
    )
    add_footer(slide, 1)


def add_use_cases_slide():
    slide = blank_slide()
    add_title(slide, "WHERE IT SHINES", "zfold が向く場面 / 向かない場面", "ZIP の置き換えというより、似たファイル群を効率よく固める用途に向きます。")
    add_card(
        slide,
        Inches(0.7),
        Inches(1.85),
        Inches(3.82),
        Inches(2.0),
        "向く 1: 小さいファイルが大量にある",
        ["HTML / JSON / 設定ファイルが何百〜何万個もある", "1 つずつ圧縮するより、まとめたほうが有利"],
        TEAL,
        TEAL_SOFT,
    )
    add_card(
        slide,
        Inches(4.76),
        Inches(1.85),
        Inches(3.82),
        Inches(2.0),
        "向く 2: ファイル同士がよく似ている",
        ["テンプレート由来の差分ファイル", "同じキー構造の JSON や設定群"],
        NAVY,
        NAVY_SOFT,
    )
    add_card(
        slide,
        Inches(8.82),
        Inches(1.85),
        Inches(3.82),
        Inches(2.0),
        "向く 3: 定期スナップショット保存",
        ["pack / list / verify / extract の流れが作りやすい", "必要なら暗号化もできる"],
        ORANGE,
        ORANGE_SOFT,
    )
    add_rect(slide, Inches(0.7), Inches(4.25), Inches(7.25), Inches(1.7), PAPER, LINE, radius=True)
    add_text(slide, Inches(0.95), Inches(4.48), Inches(6.8), Inches(0.25), "ひとことで言うと", 18, INK, True)
    add_text(slide, Inches(0.95), Inches(4.85), Inches(6.7), Inches(0.65), "「汎用交換フォーマット」よりも、\n「似た小ファイル群を効率よく固めるための専用寄りツール」", 20, NAVY, True)
    add_rect(slide, Inches(8.2), Inches(4.25), Inches(4.42), Inches(1.7), RED_SOFT, RED, radius=True)
    add_text(slide, Inches(8.45), Inches(4.48), Inches(3.9), Inches(0.25), "向きにくい例", 18, RED, True)
    add_bullets(slide, Inches(8.42), Inches(4.82), Inches(3.9), Inches(0.72), ["JPEG / PNG / MP4 が中心", "ZIP や gz など、すでに圧縮済み", "共通性の少ない大きなバイナリ"], 15, MUTED)
    add_footer(slide, 2)


def add_zip_slide():
    slide = blank_slide()
    add_title(slide, "COMPRESSION MODEL", "ZIP と zfold の圧縮の違い", "zfold は複数ファイルをチャンクにまとめてから圧縮するため、ファイル間の類似も拾えます。")
    add_zip_vs_zfold(slide)
    add_footer(slide, 3)


def add_data_slide():
    slide = blank_slide()
    add_title(slide, "GOOD FIT / BAD FIT", "圧縮率が良くなりやすいデータ、そうでもないデータ", "差が大きいのは「小さい」「似ている」「テキスト系」の条件が揃うときです。")
    add_card(
        slide,
        Inches(0.72),
        Inches(1.95),
        Inches(3.8),
        Inches(3.9),
        "得意なデータ",
        [
            "HTML / CSS / JS / TS",
            "JSON / XML / CSV / ログ",
            "INI / TOML / YAML などの設定群",
            "テンプレートから生成された類似ファイル",
        ],
        TEAL,
        TEAL_SOFT,
    )
    add_card(
        slide,
        Inches(8.8),
        Inches(1.95),
        Inches(3.8),
        Inches(3.9),
        "苦手なデータ",
        [
            "JPEG / PNG / MP4 / MP3",
            "ZIP / gz などの圧縮済みファイル",
            "内容がバラバラな巨大バイナリ",
            "共通パターンが少ないデータ",
        ],
        ORANGE,
        ORANGE_SOFT,
    )
    add_rect(slide, Inches(4.92), Inches(2.15), Inches(2.9), Inches(1.0), NAVY_SOFT, NAVY, radius=True)
    add_text(slide, Inches(5.12), Inches(2.34), Inches(2.5), Inches(0.2), "キーワード", 16, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(5.06), Inches(2.62), Inches(2.62), Inches(0.3), "小さい × 似ている × 繰り返しが多い", 13, INK, True, align=PP_ALIGN.CENTER)
    add_rect(slide, Inches(5.15), Inches(3.55), Inches(2.42), Inches(1.38), PAPER, LINE, radius=True)
    add_text(slide, Inches(5.36), Inches(3.74), Inches(2.0), Inches(0.2), "圧縮率のイメージ", 16, INK, True, align=PP_ALIGN.CENTER)
    add_rect(slide, Inches(5.45), Inches(4.15), Inches(0.44), Inches(0.45), ORANGE_SOFT, ORANGE_SOFT, radius=True)
    add_rect(slide, Inches(5.95), Inches(4.15), Inches(1.22), Inches(0.45), ORANGE, ORANGE, radius=True)
    add_text(slide, Inches(5.48), Inches(4.72), Inches(2.0), Inches(0.18), "ZIP より有利になりやすい", 13, TEAL, True, align=PP_ALIGN.CENTER)
    add_footer(slide, 4)


def add_chunk_slide():
    slide = blank_slide()
    add_title(slide, "CHUNK SIZE", "チャンクサイズの考え方", "圧縮率とメモリと部分展開のしやすさのバランスを見る設定です。")
    add_chunk_scale(slide)
    add_footer(slide, 5)


def add_dict_slide():
    slide = blank_slide()
    add_title(slide, "DICTIONARY", "辞書ファイルとは何か", "zstd の辞書圧縮を使うと、小さい類似ファイル群でさらに効くことがあります。")
    add_dictionary_flow(slide)
    add_footer(slide, 6)


def add_dict_effect_slide():
    slide = blank_slide()
    add_title(slide, "WHEN DICTIONARY HELPS", "辞書の効果が出る場面", "辞書は万能ではありません。サンプル選びが重要です。")
    add_card(
        slide,
        Inches(0.78),
        Inches(1.95),
        Inches(3.75),
        Inches(3.7),
        "効果が出やすい",
        [
            "数 KB の JSON が大量にある",
            "静的生成した HTML が大量にある",
            "設定やログに共通の見出しやキーがある",
            "毎回同じ種類のデータをアーカイブ化する",
        ],
        GREEN,
        GREEN_SOFT,
    )
    add_card(
        slide,
        Inches(8.88),
        Inches(1.95),
        Inches(3.75),
        Inches(3.7),
        "効果が小さくなりやすい",
        [
            "画像や動画が中心",
            "ZIP や gz のような圧縮済みデータ",
            "サンプルと実データの種類がズレている",
            "共通パターンの少ないバイナリ",
        ],
        RED,
        RED_SOFT,
    )
    add_rect(slide, Inches(4.92), Inches(2.1), Inches(2.85), Inches(0.95), PAPER, LINE, radius=True)
    add_text(slide, Inches(5.07), Inches(2.34), Inches(2.55), Inches(0.2), "良い辞書の条件", 17, INK, True, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(5.02), Inches(2.63), Inches(2.65), Inches(0.22), "これから圧縮したいデータに\nよく似たサンプルで作ること", 13, NAVY, True, align=PP_ALIGN.CENTER)
    add_rect(slide, Inches(4.92), Inches(3.35), Inches(2.85), Inches(1.46), NAVY_SOFT, NAVY, radius=True)
    add_text(slide, Inches(5.08), Inches(3.56), Inches(2.55), Inches(0.2), "運用上の利点", 17, NAVY, True, align=PP_ALIGN.CENTER)
    add_bullets(slide, Inches(5.08), Inches(3.88), Inches(2.5), Inches(0.7), ["辞書はアーカイブにも埋め込まれる", "展開時に別配布しなくてよい"], 13, MUTED)
    add_footer(slide, 7)


def add_workflow_slide():
    slide = blank_slide()
    add_title(slide, "WORKFLOW", "基本コマンドの流れ", "辞書なしでも始められます。効果を見てから辞書を足すのが現実的です。")
    add_commands_flow(slide)
    add_footer(slide, 8)


def add_summary_slide():
    slide = blank_slide()
    add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(7.5), NAVY, NAVY)
    add_circle(slide, Inches(10.7), Inches(0.7), Inches(1.8), TEAL, TEAL)
    add_circle(slide, Inches(11.7), Inches(1.4), Inches(0.9), ORANGE, ORANGE)
    add_text(slide, Inches(0.82), Inches(0.9), Inches(4.5), Inches(0.35), "TAKEAWAYS", 12, TEAL_SOFT, True, font="Bahnschrift")
    add_text(slide, Inches(0.82), Inches(1.32), Inches(5.0), Inches(0.7), "zfold を使う判断基準", 28, WHITE, True)
    add_card(
        slide,
        Inches(0.8),
        Inches(2.2),
        Inches(3.8),
        Inches(2.35),
        "1. ZIP と違う強み",
        ["複数ファイルをチャンク単位でまとめて圧縮", "類似した小ファイル群で圧縮率が伸びやすい"],
        TEAL,
        TEAL_SOFT,
    )
    add_card(
        slide,
        Inches(4.78),
        Inches(2.2),
        Inches(3.8),
        Inches(2.35),
        "2. 辞書はこう考える",
        ["小さくて似たテキスト系データで効きやすい", "似たサンプルから train-dict で作る"],
        ORANGE,
        ORANGE_SOFT,
    )
    add_card(
        slide,
        Inches(8.76),
        Inches(2.2),
        Inches(3.78),
        Inches(2.35),
        "3. 実運用のコツ",
        ["まず辞書なしで試す", "低メモリ環境では --threads 1 を基本線にする"],
        NAVY,
        NAVY_SOFT,
    )
    add_rect(slide, Inches(0.8), Inches(5.15), Inches(11.72), Inches(1.18), PAPER, PAPER, radius=True)
    add_text(slide, Inches(1.05), Inches(5.42), Inches(11.2), Inches(0.24), "おすすめの始め方: まずは zfold pack / verify を試し、効きそうなデータ群だけ辞書運用に進む", 19, INK, True, align=PP_ALIGN.CENTER)
    add_footer(slide, 9)


def build():
    add_intro_slide()
    add_use_cases_slide()
    add_zip_slide()
    add_data_slide()
    add_chunk_slide()
    add_dict_slide()
    add_dict_effect_slide()
    add_workflow_slide()
    add_summary_slide()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
